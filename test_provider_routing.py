# -*- coding: utf-8 -*-
"""多 provider 切换回归测试。

  - legacy_sui_service.selectBook 能切到指定旧账本 id
  - /api/login 把神象云账本打上 provider 标签;旧账本连不上时优雅降级
  - /api/book 接受 provider 并切换激活的 client
  - /api/tally 按 provider 路由(legacy 走 sui.py,shenxiang 走 client.payout/...)
  - _normalizeLegacyDetails 把旧体系的 raw 数据归一化成对账引擎可读
"""
import os
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
tmp.close()
os.environ["AUTO_SUI_DB"] = tmp.name

from contextlib import contextmanager
from fastapi.testclient import TestClient
from server import legacy_sui_service
from server import sui_service
from server.main import app, _normalizeLegacyDetails

client = TestClient(app)


@contextmanager
def patch(**kwargs):
    """轻量 monkeypatch: 把 server.x.y.z 路径的属性替换,推出 with 时还原。

    调用:patch(sui_service_createClient=..., legacy_sui_service_xxx=...) 太丑,
    改成接 kwargs={'full.dotted.path': value}。内部去掉拼接 root。
    """
    saved = {}
    try:
        for dotted, value in kwargs.items():
            parts = dotted.split(".")
            obj = sys.modules[parts[0]]
            for piece in parts[1:]:
                obj = getattr(obj, piece)
            saved[dotted] = obj
            # 重新走一遍拿到顶层,因为 obj 上面被改成了 next-level
            top = sys.modules[parts[0]]
            for piece in parts[1:-1]:
                top = getattr(top, piece)
            setattr(top, parts[-1], value)
        yield
    finally:
        for dotted, value in saved.items():
            parts = dotted.split(".")
            top = sys.modules[parts[0]]
            for piece in parts[1:-1]:
                top = getattr(top, piece)
            setattr(top, parts[-1], value)


# --------------------------------------------------------------------- #
# legacy_sui_service: 旧账本切换
#
# 账本清单已改由 books/list 动态获取(listBooks 占位项已删除),
# 这里只守住「切到指定旧账本 id」这件事。
# --------------------------------------------------------------------- #
def testSelectBookSwitchesToGivenId():
    calls = []

    class FakeSui:
        def initTallyInfo(self, bookId=None):
            calls.append(bookId)

    saved = (legacy_sui_service.buildAccountOptions,
             legacy_sui_service.buildCategoryOptions)
    legacy_sui_service.buildAccountOptions = lambda c: []
    legacy_sui_service.buildCategoryOptions = lambda c: {"income": [], "payout": []}
    try:
        legacy_sui_service.selectBook(FakeSui(), "1752598474")
    finally:
        (legacy_sui_service.buildAccountOptions,
         legacy_sui_service.buildCategoryOptions) = saved
    assert calls == ["1752598474"], calls
    print("[OK] selectBook 切到指定旧账本 id")


def testNormalizeLegacyDetailsEmpty():
    assert _normalizeLegacyDetails([]) == []
    assert _normalizeLegacyDetails(None) == []
    print("[OK] 归一化空输入安全")


def testNormalizeLegacyDetailsAlreadyShaped():
    """归一化对已是 shenxiang 形态的数据不会改写,直接透传。"""
    src = [{"sdate": "20250103", "itemAmount": 14.9, "tranType": 1,
            "sellerAcountId": "a", "buyerAcountId": "b", "tranId": "tx_001"}]
    out = _normalizeLegacyDetails(src)
    assert out[0]["sdate"] == "20250103"
    assert out[0]["itemAmount"] == 14.9
    print("[OK] shenxiang 形态透传")


def testNormalizeLegacyDetailsFuzzyMapping():
    """旧体系 detail 可能用 price/money 等字段,归一化要做宽容映射。"""
    src = [{
        "sdate": "20250103",
        "price": "12.50",
        "inoutType": 1,
        "sellerAccountId": "x",
        "buyerAccountId": "y",
        "id": "tx_9"
    }]
    out = _normalizeLegacyDetails(src)
    d = out[0]
    assert d["itemAmount"] == 12.5, "应兼容字符串金额"
    assert d["tranType"] == 1
    assert d["sellerAcountId"] == "x"
    assert d["buyerAcountId"] == "y"
    assert d["tranId"] == "tx_9"
    print("[OK] 模糊字段归一化")


def testNormalizeLegacyDetailsMissingFields():
    """字段不全也不会抛 —— 让 reconcile 当 unmatched 处理而不是 5xx。"""
    src = [{"sdate": "", "foo": "bar"}]
    out = _normalizeLegacyDetails(src)
    d = out[0]
    assert d["itemAmount"] == 0.0
    assert d["tranType"] == 0
    print("[OK] 缺字段兜底为 0,归一化不抛")


# --------------------------------------------------------------------- #
# /api/login 多 provider 行为
# --------------------------------------------------------------------- #
class FakeShenxiang:
    """够用的占位:神象云账本走 getBooks(),旧账本走 listAllBooks()。"""
    def __init__(self, books, legacyBooks=None):
        self._books = books
        self._legacyBooks = legacyBooks if legacyBooks is not None else []
    def getBooks(self):
        return list(self._books)
    def listAllBooks(self):
        return list(self._legacyBooks)
    def setBookId(self, *a, **kw): pass
    def initTallyInfo(self, *a, **kw): pass


def testLoginReturnsProviderForEachBook():
    with patch(**{
        "server.sui_service.createClient": lambda u, p: FakeShenxiang(
            [{"id": "X1", "name": "账本A"},
             {"id": "X2", "name": "账本B"}],
            # 旧账本由同一棵 token 从 books/list 取回,provider 必须是 legacy
            legacyBooks=[{"id": "1505498391", "name": "2015-", "raw": {}}]),
        "server.legacy_sui_service.createClient": lambda u, p: object(),
    }):
        r = client.post("/api/login", json={"username": "u@e.com", "password": "x"})
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["ok"] is True
    providers = [b["provider"] for b in data["books"]]
    assert providers == ["shenxiang", "shenxiang", "legacy"]
    assert data["provider"] in ("shenxiang", "legacy")
    print("[OK] /api/login 给每个账本打 provider 标签")


def testLoginLegacyFailsDoesNotBlock():
    """旧体系登录失败 -> 旧账本选项不出现,但 login 仍 200。"""
    def boom(u, p):
        raise Exception("login.sui.com 502")
    with patch(**{
        "server.sui_service.createClient": lambda u, p: FakeShenxiang([{"id": "X1", "name": "账本A"}]),
        "server.legacy_sui_service.createClient": boom,
    }):
        r = client.post("/api/login", json={"username": "u2@e.com", "password": "x"})
    assert r.status_code == 200
    data = r.json()
    assert len(data["books"]) == 1
    assert data["books"][0]["provider"] == "shenxiang"
    print("[OK] 旧体系失败时优雅降级,不阻塞登录")


def testLoginShenxiangFailsIsFatal():
    """新体系登录失败应当直接 400,不能「没有账本也登录成功」。"""
    def boom(u, p):
        raise Exception("神象云挂了")
    with patch(**{
        "server.sui_service.createClient": boom,
    }):
        r = client.post("/api/login", json={"username": "u3@e.com", "password": "x"})
    assert r.status_code == 400
    print("[OK] 新体系失败时硬拦 400")


# --------------------------------------------------------------------- #
# /api/book 接受 provider
# --------------------------------------------------------------------- #
def testSelectBookLegacyWithoutLegacyClientRejected():
    """会话没有 legacyClient(legacy 登录失败过),切到 legacy 应当 400。"""
    with patch(**{
        "server.sui_service.createClient": lambda u, p: FakeShenxiang([{"id": "X1", "name": "账本A"}]),
        "server.legacy_sui_service.createClient":
            lambda u, p: (_ for _ in ()).throw(Exception("legacy 挂")),
    }):
        r = client.post("/api/login", json={"username": "u4@e.com", "password": "x"})
    sid = r.json()["sid"]
    r = client.post("/api/book",
                    json={"sid": sid, "bookId": legacy_sui_service.LEGACY_BOOK_ID,
                          "provider": "legacy"})
    assert r.status_code == 400
    print("[OK] legacy 不可用时拒绝切换")


def testSelectBookUnknownProviderRejected():
    """provider 不在白名单返回 400。"""
    with patch(**{
        "server.sui_service.createClient": lambda u, p: FakeShenxiang([{"id": "X1", "name": "A"}]),
        "server.legacy_sui_service.createClient":
            lambda u, p: (_ for _ in ()).throw(Exception("x")),
    }):
        r = client.post("/api/login", json={"username": "u5@e.com", "password": "x"})
    sid = r.json()["sid"]
    r = client.post("/api/book", json={"sid": sid, "bookId": "X1",
                                       "provider": "something_bad"})
    assert r.status_code == 400
    print("[OK] provider 不在白名单被 400")


# --------------------------------------------------------------------- #
# /api/tally 按 provider 路由
# --------------------------------------------------------------------- #
class CountingLegacy:
    """能记录每次被调用的服务对象。

    注:payout/income/transfer 现在走 legacy_sui_service._checkResult —— 它会
    校验响应(状态码/HTML/JSON),所以这里要给一个长得像成功响应的对象,否则
    测试会在校验阶段就挂掉,看不到「是否被调用」。
    """
    def __init__(self):
        self.calls = []
    def initTallyInfo(self, bookId=None):
        self.calls.append(("initTallyInfo", bookId))
    # selectBook 之后会读账户/分类,给个空壳即可
    def getAccounts(self):
        return []
    def getCategories(self):
        return {"income": [], "payout": []}
    @staticmethod
    def _ok():
        return type("R", (), {"status_code": 200,
                              "text": "{id:{id:134923109120449},budget:0,price:0.01}"})()
    def payout(self, *a, **kw):
        self.calls.append(("payout", a, kw))
        return self._ok()
    def income(self, *a, **kw):
        self.calls.append(("income", a, kw))
        return self._ok()
    def transfer(self, *a, **kw):
        self.calls.append(("transfer", a, kw))
        return self._ok()


def testTallyRoutesToLegacy():
    """legacy 路径下 payout 应当调 CountingLegacy.payout 而非 shenxiang client。"""
    fake_lex = CountingLegacy()
    fake_shx = FakeShenxiang([{"id": "X1", "name": "A"}])
    with patch(**{
        "server.sui_service.createClient": lambda u, p: fake_shx,
        "server.legacy_sui_service.createClient": lambda u, p: fake_lex,
    }):
        r = client.post("/api/login", json={"username": "u6@e.com", "password": "x"})
    sid = r.json()["sid"]
    # 切到 legacy
    r = client.post("/api/book",
                    json={"sid": sid,
                          "bookId": legacy_sui_service.LEGACY_BOOK_ID,
                          "provider": "legacy"})
    assert r.status_code == 200, r.text
    # 准备 pending 然后调 tally
    from server import session_store as ss
    sess = ss.get(sid)
    sess["pending"] = {
        "source": "upload", "bankno": "622848****1211",
        "startDate": "20260101", "endDate": "20260131",
        "details": [{"amount": 14.9, "date": "20260105", "time": "0800",
                     "transType": "payout", "memo": "测试"}],
        "suiid": legacy_sui_service.LEGACY_BOOK_ID
    }
    r = client.post("/api/tally", json={
        "sid": sid, "suiid": "17330926177",
        "index": 0, "op": "payout",
        "catid": "51160931276", "memo": "auto"
    })
    assert r.status_code == 200, r.text
    # 选账本会先触发 initTallyInfo(切换旧账本),之后 tally 才产生记账调用
    kinds = [c[0] for c in fake_lex.calls]
    assert "initTallyInfo" in kinds, "选旧账本应触发账本切换"
    assert "payout" in kinds, "legacy 应该收到 payout 调用"
    op_name, args, kw = next(c for c in fake_lex.calls if c[0] == "payout")
    assert args[1] == 14.9
    assert kw["memo"] == "auto"
    print("[OK] /api/tally 走 legacy 服务")


def testUploadDropsSuiidWhenItDoesNotBelongToLegacyBook():
    """上传账单:神象云映射的 suiid 在旧账本里不存在时,必须清空让用户手选。

    之前 bug:直接把 conf.json 里的神象云 id 给前端,导致对账查空账户、记账被
    服务端拒绝却显示成功。
    """
    import openpyxl, tempfile
    from unittest.mock import patch

    # 准备一份合法的小 xlsx
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet0"
    ws["A2"] = "账户：622848****1211 起始日期：20260101 截止日期：20260131"
    ws.cell(row=4, column=1, value="2026-01-05")
    ws.cell(row=4, column=2, value="08:00:00")
    ws.cell(row=4, column=3, value="-14.90")
    ws.cell(row=4, column=4, value=100.0)
    for col in range(5, 12):
        ws.cell(row=4, column=col, value="x")
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    wb.save(tmp.name)
    with open(tmp.name, "rb") as f:
        bytes_ = f.read()

    fake_lex = CountingLegacy()
    fake_shx = FakeShenxiang([{"id": "X1", "name": "A"}])
    sid = None
    with patch("server.sui_service.createClient", lambda u, p: fake_shx), \
         patch("server.legacy_sui_service.createClient", lambda u, p: fake_lex):
        r = client.post("/api/login", json={"username": "u-cross@e.com", "password": "x"})
        assert r.status_code == 200
        sid = r.json()["sid"]
        # 切到旧账本
        r = client.post("/api/book", json={"sid": sid,
                                            "bookId": legacy_sui_service.LEGACY_BOOK_ID,
                                            "provider": "legacy"})
        assert r.status_code == 200
        # 先把神象云 id 塞进映射表 —— 用户真实环境里就是 conf.json 导入的
        r = client.post("/api/mapping", json={"sid": sid, "provider": "shenxiang",
                                              "accounts": [{"bankno": "622848****1211",
                                                             "suiid": "1183357936638279681",
                                                             "type": "abc_debit"}]})
        assert r.status_code == 200, r.text
        r = client.post("/api/bills/upload",
                        data={"sid": sid, "bankType": "abc"},
                        files={"file": ("abc.xlsx", bytes_,
                                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["suiid"] == "", "shenxiang id 在旧账本下不应被返回,期望空,实际: %s" % body["suiid"]
    print("[OK] 上传账单把跨 provider 的 suiid 清空")


def testLegacyTallyReportsFailureOnHtmlErrorPage():
    """sui.py 返回 HTML 错误页(账户 id 错、参数非法)时,/api/tally 必须 400。

    之前 bug:sui.payout 静默忽略,服务端仍然报「记账成功」。
    """
    from unittest.mock import patch
    from server import session_store as ss

    class FailingLegacy(CountingLegacy):
        def payout(self, *a, **kw):
            self.calls.append(("payout", a, kw))
            return type("R", (), {"status_code": 200,
                                  "text": "<html><body>出错啦 服务器开小差了</body></html>"})()

    fake_lex = FailingLegacy()
    fake_shx = FakeShenxiang([{"id": "X1", "name": "A"}])
    with patch("server.sui_service.createClient", lambda u, p: fake_shx), \
         patch("server.legacy_sui_service.createClient", lambda u, p: fake_lex):
        r = client.post("/api/login", json={"username": "u-fail@e.com", "password": "x"})
        sid = r.json()["sid"]
        client.post("/api/book", json={"sid": sid,
                                       "bookId": legacy_sui_service.LEGACY_BOOK_ID,
                                       "provider": "legacy"})
        ss.get(sid)["pending"] = {
            "source": "upload", "bankno": "622848****1211",
            "startDate": "20260101", "endDate": "20260131",
            "details": [{"amount": 14.9, "date": "20260105", "time": "0800",
                         "transType": "payout", "memo": "测试"}],
            "suiid": legacy_sui_service.LEGACY_BOOK_ID,
        }
        r = client.post("/api/tally", json={
            "sid": sid, "suiid": "17330926177", "index": 0,
            "op": "payout", "catid": "51160931276", "memo": "auto"
        })
    assert r.status_code == 400, r.text
    assert "被随手记拒绝" in r.json().get("detail", ""), r.text
    print("[OK] /api/tally 把 HTML 错误页当成失败上报")


if __name__ == "__main__":
    testSelectBookSwitchesToGivenId()
    testNormalizeLegacyDetailsEmpty()
    testNormalizeLegacyDetailsAlreadyShaped()
    testNormalizeLegacyDetailsFuzzyMapping()
    testNormalizeLegacyDetailsMissingFields()
    testLoginReturnsProviderForEachBook()
    testLoginLegacyFailsDoesNotBlock()
    testLoginShenxiangFailsIsFatal()
    testSelectBookLegacyWithoutLegacyClientRejected()
    testSelectBookUnknownProviderRejected()
    testTallyRoutesToLegacy()
    testUploadDropsSuiidWhenItDoesNotBelongToLegacyBook()
    testLegacyTallyReportsFailureOnHtmlErrorPage()
    print("\n多 provider 切换测试通过。")

