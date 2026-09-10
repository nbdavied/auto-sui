# -*- coding: utf-8 -*-
"""绕开 HTTP 登录,直接调用 server 的对账 / 记账 endpoint,复现真实链路。

用例取自线上真实数据(账本 1505498391,账户 17330926177 农业银行1211):
  - 20260807 支出 10000.00 —— 账本里已有,应判定 matched
  - 20260909 支出 1.23    —— 账本里没有,应判定 manual/auto
然后对未匹配的那条发起记账,再重新对账看是否变成 matched。
"""
import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl  # noqa: E402

from server import main as M  # noqa: E402
from server import readers, store  # noqa: E402
from server.reconcile import transDate  # noqa: E402

BOOK = "1505498391"
ACCOUNT = "17330926177"
START, END = "20260101", "20260910"

conf = json.load(open("conf.json", encoding="utf-8"))


def buildBill():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet0"
    ws["A2"] = "账户：622848****1211 起始日期：%s 截止日期：%s" % (START, END)
    rows = [
        ("2026-08-07", "12:00:00", "-10000.00", "对手A", "111", "用途A", "备注A"),
        ("2026-09-09", "08:30:00", "-1.23", "对手B", "222", "用途B", "备注B"),
    ]
    for i, (d, t, amt, name, no, usage, memo) in enumerate(rows):
        ws.cell(row=4 + i, column=1, value=d)
        ws.cell(row=4 + i, column=2, value=t)
        ws.cell(row=4 + i, column=3, value=amt)
        ws.cell(row=4 + i, column=4, value=100.0)
        ws.cell(row=4 + i, column=5, value=name)
        ws.cell(row=4 + i, column=6, value=no)
        ws.cell(row=4 + i, column=7, value="")
        ws.cell(row=4 + i, column=8, value="")
        ws.cell(row=4 + i, column=10, value=usage)
        ws.cell(row=4 + i, column=11, value=memo)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


sid = M.sessions.create()
s = M.sessions.get(sid)
user = store.getOrCreateUser(conf["username"])
s["username"] = conf["username"]
s["password"] = conf["password"]
s["userId"] = user["id"]
s["legacyClient"] = M.legacy_sui_service.createClient(conf["username"], conf["password"])
s["provider"] = M.PROVIDER_LEGACY

print("== 切换账本 ==")
res = M.selectBook(M.BookBody(sid=sid, bookId=BOOK, provider=M.PROVIDER_LEGACY))
print("账户数", len(res["accounts"]), "支出一级", len(res["categories"]["payout"]),
      "provider", res["provider"])

print("== 账户 / 分类 ==")
acc = M.listAccounts(sid)
print("账户数", len(acc["accounts"]))
cats = M.listCategories(sid)
print("支出一级", len(cats["categories"]["payout"]), "收入一级", len(cats["categories"]["income"]))

print("== 解析账单 ==")
data = readers.parseUpload(user["id"], buildBill(), "abc.xlsx", "abc")
print("bankno", data["bankno"], "suiid", data["suiid"], "范围", data["startDate"], data["endDate"])
s["pending"] = data
print(taskRange := (transDate(data["startDate"]), transDate(data["endDate"])))

print("== 对账 ==")
res = M.reconcileBills(M.ReconcileBody(sid=sid, suiid=ACCOUNT))
for i, it in enumerate(res["items"]):
    print(i, it["date"], it["amount"], it["transType"], "->", it["status"],
          "| matched:", json.dumps(it.get("matched"), ensure_ascii=False) if it.get("matched") else None)

print("== 对未匹配条目记账 ==")
target = None
for i, it in enumerate(res["items"]):
    if it["status"] != "matched":
        target = i
        break
if target is None:
    print("没有可记账的条目")
else:
    catid = cats["categories"]["payout"][0]["children"][0]["id"]
    print("index", target, "catid", catid)
    try:
        r = M.tally(M.TallyBody(sid=sid, suiid=ACCOUNT, index=target, op="payout",
                                catid=catid, memo="auto-sui e2e"))
        print("记账返回", r)
    except Exception as e:
        print("记账异常:", type(e).__name__, getattr(e, "detail", e))

    print("== 重新对账 ==")
    res2 = M.reconcileBills(M.ReconcileBody(sid=sid, suiid=ACCOUNT))
    for i, it in enumerate(res2["items"]):
        print(i, it["date"], it["amount"], "->", it["status"])
